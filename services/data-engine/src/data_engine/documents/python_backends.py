"""Python document back-ends. Each import is optional; `HAS` records what is installed."""

from __future__ import annotations

import importlib
import os
import shutil
from io import StringIO
from pathlib import Path

import pandas as pd

from .. import parallel


def _has(module: str) -> bool:
    try:
        importlib.import_module(module)
        return True
    except Exception:
        return False


HAS = {name: _has(mod) for name, mod in {
    "pdfplumber": "pdfplumber", "pypdf": "pypdf", "pypdfium2": "pypdfium2", "docx": "docx", "pptx": "pptx",
    "bs4": "bs4", "lxml": "lxml", "openpyxl": "openpyxl", "tika": "tika", "textract": "textract",
}.items()}
def _tesseract_cmd() -> str | None:
    """Tesseract binary: TESSERACT_CMD, PATH, or the default Windows install location."""
    for candidate in (os.environ.get("TESSERACT_CMD"), shutil.which("tesseract"),
                      r"C:\Program Files\Tesseract-OCR\tesseract.exe"):
        if candidate and Path(candidate).exists():
            return candidate
    return None


HAS["pytesseract"] = _has("pytesseract") and _tesseract_cmd() is not None
if HAS["pytesseract"]:
    import pytesseract as _pt
    _pt.pytesseract.tesseract_cmd = _tesseract_cmd()


def available() -> dict[str, bool]:
    return dict(HAS)


def default_ocr_lang() -> str:
    """English + Spanish when both language packs are installed (breeders write both)."""
    env = os.environ.get("DATA_ENGINE_OCR_LANG")
    if env:
        return env
    if not HAS["pytesseract"]:
        return "eng"
    import pytesseract
    try:
        langs = set(pytesseract.get_languages(config=""))
    except Exception:
        return "eng"
    return "+".join(lang for lang in ("eng", "spa") if lang in langs) or "eng"


# ----------------------------------------------------------------------------- PDF
def pdf_text_pdfplumber(path: Path) -> list[str]:
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return [(page.extract_text() or "") for page in pdf.pages]


def pdf_text_pypdf(path: Path) -> list[str]:
    from pypdf import PdfReader
    return [(page.extract_text() or "") for page in PdfReader(str(path)).pages]


def pdf_tables(path: Path) -> list[dict]:
    """Tables with their structure (rows/columns), one DataFrame per table, header row kept as row 0."""
    import pdfplumber
    out = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            for rows in page.extract_tables() or []:
                if rows and len(rows) > 1 and len(rows[0]) > 1:
                    out.append({"page": i + 1, "frame": pd.DataFrame(rows), "source": "pdfplumber"})
    return out


def _render_page(path: Path, index: int, dpi: int = 300):
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(path))
    try:
        return pdf[index].render(scale=dpi / 72.0).to_pil()
    finally:
        pdf.close()


def _ocr_pil(image, lang: str) -> tuple[str, float | None]:
    """Text and mean word confidence (Tesseract's 0-100 scale, -1 entries ignored)."""
    import pytesseract
    data = pytesseract.image_to_data(image, lang=lang, output_type=pytesseract.Output.DICT)
    words, confs = [], []
    line_key, lines = None, []
    for i, word in enumerate(data["text"]):
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        if key != line_key and words:
            lines.append(" ".join(words))
            words = []
        line_key = key
        if word.strip():
            words.append(word)
            conf = float(data["conf"][i])
            if conf >= 0:
                confs.append(conf)
    if words:
        lines.append(" ".join(words))
    return "\n".join(lines), (round(sum(confs) / len(confs), 2) if confs else None)


def pdf_ocr_pages(path: Path, pages: list[int], lang: str, dpi: int = 300) -> tuple[list[str], float | None, dict]:
    """OCR the given pages in parallel. Tesseract runs out of process, so threads parallelise it."""
    decision = parallel.plan_tasks(costs=[1.5] * len(pages), gil_share=0.05,
                                   peak_memory=[parallel.task_memory(0, "ocr")] * len(pages))
    if decision.workers > 1:
        os.environ.setdefault("OMP_THREAD_LIMIT", "1")  # one core per Tesseract process: no oversubscription

    def work(i: int):
        return _ocr_pil(_render_page(path, i, dpi), lang)

    results = parallel.run(work, pages, decision)
    confs = [c for _, c in results if c is not None]
    return [t for t, _ in results], (round(sum(confs) / len(confs), 2) if confs else None), decision.as_dict()


def ocr_image(path: Path, lang: str) -> tuple[str, float | None]:
    from PIL import Image
    with Image.open(path) as img:
        return _ocr_pil(img.convert("RGB"), lang)


# ----------------------------------------------------------------------------- Office
def docx_extract(path: Path) -> tuple[list[str], list[dict]]:
    import docx
    d = docx.Document(str(path))
    text = "\n".join(p.text for p in d.paragraphs)
    tables = [{"page": 1, "frame": pd.DataFrame([[c.text for c in row.cells] for row in t.rows]), "source": "docx"}
              for t in d.tables if len(t.rows) > 1]
    return [text], tables


def pptx_extract(path: Path) -> tuple[list[str], list[dict]]:
    from pptx import Presentation
    pages, tables = [], []
    for i, slide in enumerate(Presentation(str(path)).slides):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                if len(rows) > 1:
                    tables.append({"page": i + 1, "frame": pd.DataFrame(rows), "source": "pptx"})
        pages.append("\n".join(parts))
    return pages, tables


# ----------------------------------------------------------------------------- HTML
def html_extract(path: Path) -> tuple[list[str], list[dict]]:
    """Visible text + <table>s. The same parsing toolbox as web scraping, applied to local files."""
    raw = path.read_text(encoding="utf-8", errors="replace")
    tables = []
    if HAS["lxml"]:
        try:
            tables = [{"page": 1, "frame": t, "source": "html"} for t in pd.read_html(StringIO(raw), header=None)]
        except ValueError:
            tables = []  # no <table> in the page
    if HAS["bs4"]:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(raw, "lxml" if HAS["lxml"] else "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        if not tables:
            for t in soup.find_all("table"):
                rows = [[c.get_text(strip=True) for c in tr.find_all(["td", "th"])] for tr in t.find_all("tr")]
                if len(rows) > 1:
                    tables.append({"page": 1, "frame": pd.DataFrame(rows), "source": "html"})
        text = soup.get_text("\n", strip=True)
    else:
        from html.parser import HTMLParser

        class _Text(HTMLParser):
            def __init__(self):
                super().__init__()
                self.parts, self.skip = [], False

            def handle_starttag(self, tag, attrs):
                self.skip = tag in ("script", "style")

            def handle_data(self, data):
                if not self.skip and data.strip():
                    self.parts.append(data.strip())

        parser = _Text()
        parser.feed(raw)
        text = "\n".join(parser.parts)
    return [text], tables


# ----------------------------------------------------------------------------- Generic
def tika_text(path: Path) -> str:
    """Apache Tika (Java). Needs Java; set TIKA_SERVER_JAR for offline use."""
    from tika import parser
    return (parser.from_file(str(path)) or {}).get("content") or ""


def textract_text(path: Path) -> str:
    """textract: broad but unmaintained (old pinned dependencies); last-resort fallback."""
    import textract
    return textract.process(str(path)).decode("utf-8", errors="replace")
