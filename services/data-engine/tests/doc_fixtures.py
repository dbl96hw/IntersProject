"""Generate small test documents on the fly (no binary fixtures in the repo).

Each generator returns None when the library it needs is missing, so the
corresponding test is skipped instead of failing on a lighter install.
"""

from __future__ import annotations

from pathlib import Path

GENOMICS_HEADER = ["GENOMIC_SAMPLE_GUID", "MATERIAL_GUID", "MARKER_DROUGHT_TOLERANCE", "MARKER_DISEASE_RESISTANCE",
                   "MARKER_YIELD_POTENTIAL", "MARKER_MATURITY", "GENOMIC_BREEDING_VALUE", "QC_CALL_RATE_PCT"]
GENOMICS_ROW = ["23C6D4C2-C046-2378-0000-000000000001", "1FC9E916-899C-6A96-0000-000000000001", "FAVOURABLE",
                "RESISTANT", "HIGH", "MID", "111.5", "98.2"]
NOTE = "Field note: SYN-MZ-00012 in SYN-TR-0007 showed yield 9.4 t/ha and disease score 3.5."
NOTE_LINES = ["Field note: SYN-MZ-00012 in SYN-TR-0007", "showed yield 9.4 t/ha and disease score 3.5."]


def text_image(path: Path, lines: list[str]):
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return None
    img = Image.new("RGB", (1700, 400), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 40)
    except OSError:
        try:
            font = ImageFont.load_default(size=40)
        except TypeError:
            font = ImageFont.load_default()
    for i, line in enumerate(lines):
        draw.text((40, 40 + 70 * i), line, fill="black", font=font)
    img.save(path)
    return path


def scanned_pdf(path: Path):
    """An image-only PDF (no text layer): what a scanner produces."""
    png = text_image(path.with_suffix(".png"), NOTE_LINES)
    if png is None:
        return None
    from PIL import Image
    Image.open(png).convert("RGB").save(path, "PDF", resolution=150)
    return path


def native_pdf_with_table(path: Path):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape, A3
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        return None
    table = Table([GENOMICS_HEADER, GENOMICS_ROW], repeatRows=1)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("FONTSIZE", (0, 0), (-1, -1), 5)]))
    SimpleDocTemplate(str(path), pagesize=landscape(A3)).build(
        [Paragraph(NOTE, getSampleStyleSheet()["Normal"]), table])
    return path


def docx_with_table(path: Path):
    try:
        import docx
    except ImportError:
        return None
    d = docx.Document()
    d.add_paragraph(NOTE)
    t = d.add_table(rows=2, cols=len(GENOMICS_HEADER))
    for j, (h, v) in enumerate(zip(GENOMICS_HEADER, GENOMICS_ROW)):
        t.cell(0, j).text, t.cell(1, j).text = h, v
    d.save(path)
    return path


def html_with_table(path: Path):
    head = "".join(f"<th>{h}</th>" for h in GENOMICS_HEADER)
    row = "".join(f"<td>{v}</td>" for v in GENOMICS_ROW)
    path.write_text(f"<html><body><p>{NOTE}</p><table><tr>{head}</tr><tr>{row}</tr></table>"
                    f"<script>var x = 'ignored';</script></body></html>", encoding="utf-8")
    return path


def pptx_note(path: Path):
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        return None
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Advancement meeting"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1))
    box.text_frame.text = NOTE
    prs.save(path)
    return path
